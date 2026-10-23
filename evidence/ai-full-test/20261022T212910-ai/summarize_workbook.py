import json, hashlib
from openpyxl import load_workbook
path=r'D:\BookSys\Bán vé sự kiện có sơ đồ ghế.xlsx'; outdir=r'D:\Merge_2\Codegym\evidence\ai-full-test\20261022T212910-ai'
wb=load_workbook(path,read_only=True,data_only=True)
def records(name):
 ws=wb[name]; rows=list(ws.iter_rows(values_only=True)); hdr=[str(x).strip() if x is not None else '' for x in rows[0]]
 return ws,hdr,[dict(zip(hdr,row)) for row in rows[1:] if any(v is not None for v in row)]
inv={}
for name in ['Backlog','Tasks','Epics','Sprints','Rủi ro','DoD-DoR']:
 ws,hdr,rs=records(name)
 inv[name]={'rows':len(rs),'headers':hdr}
 if name=='Backlog':
  inv[name]['entries']=[{'id':r.get('ID'),'type':r.get('Type'),'title':r.get('Title'),'tier':r.get('Tier'),'priority':r.get('Priority'),'sprint':r.get('Sprint'),'status':r.get('Status'),'source_row':i+2,'source_cells':f'Backlog!A{i+2}:O{i+2}'} for i,r in enumerate(rs)]
  inv[name]['tier_counts']={str(k):sum(r.get('Tier')==k for r in rs) for k in sorted(set(r.get('Tier') for r in rs),key=str)}
  inv[name]['priority_counts']={str(k):sum(r.get('Priority')==k for r in rs) for k in sorted(set(r.get('Priority') for r in rs),key=str)}
  inv[name]['status_counts']={str(k):sum(r.get('Status')==k for r in rs) for k in sorted(set(r.get('Status') for r in rs),key=str)}
 if name=='Tasks': inv[name]['entries']=[{'id':r.get('ID'),'parent':r.get('Parent'),'title':r.get('Title'),'sprint':r.get('Sprint'),'status':r.get('Status'),'source_row':i+2,'source_cells':f'Tasks!A{i+2}:J{i+2}'} for i,r in enumerate(rs)]
 if name=='Epics': inv[name]['entries']=[{'id':r.get('ID'),'title':r.get('Title'),'tier':r.get('Tier'),'priority':r.get('Priority'),'status':r.get('Status'),'source_row':i+2,'source_cells':f'Epics!A{i+2}:J{i+2}'} for i,r in enumerate(rs)]
with open(outdir+r'\workbook-baseline-summary.json','w',encoding='utf-8') as f: json.dump(inv,f,ensure_ascii=False,indent=2,default=str)
print(json.dumps({k:{x:y for x,y in v.items() if x!='entries'} for k,v in inv.items()},ensure_ascii=False,indent=2))
