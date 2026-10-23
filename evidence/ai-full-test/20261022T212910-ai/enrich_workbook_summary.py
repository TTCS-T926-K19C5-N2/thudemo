import json
from openpyxl import load_workbook
p=r'D:\BookSys\Bán vé sự kiện có sơ đồ ghế.xlsx'; outp=r'D:\Merge_2\Codegym\evidence\ai-full-test\20261022T212910-ai\workbook-baseline-summary.json'; wb=load_workbook(p,read_only=True,data_only=True); d=json.load(open(outp,encoding='utf-8'))
def rows(name):
 ws=wb[name]; vals=list(ws.iter_rows(values_only=True)); hdr=[str(x).strip() if x is not None else '' for x in vals[0]]; return ws,hdr,[dict(zip(hdr,r)) for r in vals[1:] if any(v is not None for v in r)]
ws,hdr,rs=rows('Thông tin'); d['Project Goal']={'source_row':'Thông tin!A1:B14','entries':[{'row':i+2,'field':r.get('Trường'),'value':r.get('Giá trị')} for i,r in enumerate(rs)]}
ws,hdr,rs=rows('Sprints'); d['Sprints']['entries']=[{'row':i+2,'source_cells':f'Sprints!A{i+2}:F{i+2}',**r} for i,r in enumerate(rs)]
ws,hdr,rs=rows('DoD-DoR'); d['DoD-DoR']['entries']=[{'row':i+2,'source_cells':f'DoD-DoR!A{i+2}:C{i+2}',**r} for i,r in enumerate(rs)]
ws,hdr,rs=rows('Rủi ro'); d['Rủi ro']['entries']=[{'row':i+2,'source_cells':f'Rủi ro!A{i+2}:F{i+2}',**r} for i,r in enumerate(rs)]
with open(outp,'w',encoding='utf-8') as f: json.dump(d,f,ensure_ascii=False,indent=2,default=str)
