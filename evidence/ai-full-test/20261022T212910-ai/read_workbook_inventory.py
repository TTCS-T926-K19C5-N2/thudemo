import json, hashlib, os
from openpyxl import load_workbook
path = r'D:\BookSys\Bán vé sự kiện có sơ đồ ghế.xlsx'
wb = load_workbook(path, read_only=True, data_only=False)
out = {'source': path, 'sha256': hashlib.sha256(open(path,'rb').read()).hexdigest(), 'sheets': []}
for ws in wb.worksheets:
    rows=[]
    for row in ws.iter_rows(values_only=True):
        vals=[v for v in row]
        if any(v is not None for v in vals): rows.append(vals)
    out['sheets'].append({'name':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'nonempty_rows':len(rows),'rows':rows})
with open(r'D:\Merge_2\Codegym\evidence\ai-full-test\20261022T212910-ai\workbook-inventory.json','w',encoding='utf-8') as f: json.dump(out,f,ensure_ascii=False,default=str,indent=2)
print(json.dumps([{'name':s['name'],'rows':s['nonempty_rows'],'columns':s['max_column'],'headers':s['rows'][0] if s['rows'] else []} for s in out['sheets']],ensure_ascii=False,indent=2))
